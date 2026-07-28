from pathlib import Path
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
for p in map(Path, sys.argv[1:]):
    print(f"\n=== {p.name} ===")
    wb_formula = load_workbook(p, data_only=False, read_only=False)
    wb_values = load_workbook(p, data_only=True, read_only=False)
    print("sheets", wb_formula.sheetnames)
    for name, area in [("Settings", "A1:B12"), ("Raw Data", "A1:J15"), ("Evaluation", "A1:X15"), ("Metrics", "A1:Q7")]:
        wsf = wb_formula[name]
        wsv = wb_values[name]
        print(f"-- {name} {area} dim={wsf.max_row}x{wsf.max_column}")
        for row in wsf[area]:
            parts = []
            for c in row:
                v = c.value
                cached = wsv[c.coordinate].value
                if v is not None or cached is not None:
                    parts.append(f"{c.coordinate}={v!r}[{c.data_type}]=>{cached!r}")
            if parts:
                print(" | ".join(parts))
    for ws in wb_formula.worksheets:
        print(f"charts {ws.title}: {len(ws._charts)}")
