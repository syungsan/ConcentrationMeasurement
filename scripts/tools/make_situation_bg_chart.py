from __future__ import annotations

import argparse
import math
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = PROJECT_ROOT / "outputs" / "agg_timeseries_sec.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "outputs" / "agg_timeseries_sec_situation_bg.xlsx"


SITUATION_COLORS = {
    "聞く": (211, 235, 255),
    "書く": (255, 239, 179),
    "話し合う": (222, 241, 206),
    "AB": (255, 210, 210),
    None: (235, 235, 235),
    "": (235, 235, 235),
}
LINE_COLOR = (30, 64, 175)
GRID_COLOR = (215, 221, 230)
TEXT_COLOR = (35, 45, 60)


def load_font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\meiryob.ttc") if bold else Path(r"C:\Windows\Fonts\meiryo.ttc"),
        Path(r"C:\Windows\Fonts\YuGothB.ttc") if bold else Path(r"C:\Windows\Fonts\YuGothR.ttc"),
        Path(r"C:\Windows\Fonts\msgothic.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def draw_text(draw: ImageDraw.ImageDraw, xy, value: str, font, fill=TEXT_COLOR, anchor=None):
    try:
        draw.text(xy, value, font=font, fill=fill, anchor=anchor)
    except UnicodeEncodeError:
        draw.text(xy, value.encode("ascii", "ignore").decode("ascii"), font=font, fill=fill, anchor=anchor)


def nice_ticks(vmin: float, vmax: float, count: int = 7):
    if vmax <= vmin:
        return [vmin]
    raw = (vmax - vmin) / max(1, count - 1)
    exp = math.floor(math.log10(raw))
    base = raw / (10**exp)
    step_base = 1 if base <= 1 else 2 if base <= 2 else 5 if base <= 5 else 10
    step = step_base * (10**exp)
    start = math.floor(vmin / step) * step
    ticks = []
    v = start
    while v <= vmax + step * 0.5:
        if v >= vmin - step * 0.5:
            ticks.append(round(v, 6))
        v += step
    return ticks


def collect_rows(ws, *, situation_col: str, x_col: str, score_col: str):
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    idx = {name: i + 1 for i, name in enumerate(headers)}
    required = [situation_col, x_col, score_col]
    missing = [name for name in required if name not in idx]
    if missing:
        raise RuntimeError(f"必要な列が見つかりません: {', '.join(missing)}")

    rows = []
    for r in range(2, ws.max_row + 1):
        x = ws.cell(r, idx[x_col]).value
        score = ws.cell(r, idx[score_col]).value
        if x is None or score is None:
            continue
        rows.append((float(x), ws.cell(r, idx[situation_col]).value, float(score)))
    if not rows:
        raise RuntimeError("グラフ化できる時系列データが見つかりません。")
    return rows


def build_situation_spans(rows):
    xs = [r[0] for r in rows]
    dt_candidates = [b - a for a, b in zip(xs, xs[1:]) if b > a]
    dt = sorted(dt_candidates)[len(dt_candidates) // 2] if dt_candidates else 1.0

    spans = []
    start = rows[0][0]
    current_situation = rows[0][1]
    for i in range(1, len(rows)):
        if rows[i][1] != current_situation:
            spans.append((start, rows[i][0], current_situation))
            start = rows[i][0]
            current_situation = rows[i][1]
    spans.append((start, rows[-1][0] + dt, current_situation))
    return spans, dt


def make_chart_png(
    rows,
    spans,
    png_path: Path,
    *,
    title: str,
    x_label: str,
    y_label: str,
    width: int,
    height: int,
):
    xs = [r[0] for r in rows]
    scores = [r[2] for r in rows]
    x_min, x_max = min(xs), max(s[1] for s in spans)
    y_min, y_max = 1.0, 7.0
    if min(scores) < y_min or max(scores) > y_max:
        pad = max(0.25, (max(scores) - min(scores)) * 0.08)
        y_min = min(y_min, math.floor((min(scores) - pad) * 2) / 2)
        y_max = max(y_max, math.ceil((max(scores) + pad) * 2) / 2)

    left, right, top, bottom = 95, 45, 85, 105
    plot_w = width - left - right
    plot_h = height - top - bottom

    img = Image.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    font_title = load_font(30, bold=True)
    font_label = load_font(18)
    font_small = load_font(15)

    def xmap(x):
        return left + (x - x_min) / (x_max - x_min) * plot_w

    def ymap(y):
        return top + (y_max - y) / (y_max - y_min) * plot_h

    draw_text(draw, (left, 28), title, font_title)

    for s0, s1, situation in spans:
        color = SITUATION_COLORS.get(situation, (235, 235, 235))
        draw.rectangle([xmap(s0), top, xmap(s1), top + plot_h], fill=color)

    for y in nice_ticks(y_min, y_max, 7):
        yy = ymap(y)
        draw.line([left, yy, left + plot_w, yy], fill=GRID_COLOR, width=1)
        draw_text(draw, (left - 12, yy), f"{y:g}", font_small, anchor="rm")

    for i in range(11):
        x = x_min + (x_max - x_min) * i / 10
        xx = xmap(x)
        draw.line([xx, top + plot_h, xx, top + plot_h + 6], fill=TEXT_COLOR, width=1)
        draw_text(draw, (xx, top + plot_h + 12), f"{x:.0f}", font_small, anchor="ma")

    draw.rectangle([left, top, left + plot_w, top + plot_h], outline=(80, 95, 120), width=2)
    draw_text(draw, (left + plot_w / 2, height - 38), x_label, font_label, anchor="mm")
    draw_text(draw, (left + 4, top - 28), y_label, font_label)

    points = [(xmap(x), ymap(score)) for x, _situation, score in rows]
    if len(points) >= 2:
        draw.line(points, fill=LINE_COLOR, width=4, joint="curve")

    legend_items = []
    seen = {situation for _s0, _s1, situation in spans}
    for label in ("聞く", "書く", "話し合う", "AB"):
        if label in seen:
            legend_items.append((label, SITUATION_COLORS[label]))
    legend_x, legend_y = left + plot_w - min(610, 145 * max(1, len(legend_items))), 35
    for j, (label, color) in enumerate(legend_items):
        x = legend_x + j * 145
        draw.rectangle([x, legend_y, x + 28, legend_y + 18], fill=color, outline=(120, 120, 120))
        draw_text(draw, (x + 36, legend_y - 2), label, font_small)
    draw.line([legend_x, legend_y + 42, legend_x + 35, legend_y + 42], fill=LINE_COLOR, width=4)
    draw_text(draw, (legend_x + 43, legend_y + 32), y_label, font_small)

    png_path.parent.mkdir(parents=True, exist_ok=True)
    img.save(png_path, quality=95)


def add_chart_sheet(wb, png_path: Path, spans, *, sheet_name: str, title: str):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    chart_ws = wb.create_sheet(sheet_name, 0)
    chart_ws.sheet_view.showGridLines = False
    chart_ws["A1"] = title
    chart_ws["A1"].font = Font(bold=True, size=16)

    xl_img = XLImage(str(png_path))
    xl_img.anchor = "A3"
    chart_ws.add_image(xl_img)

    start_row = 45
    headers = ["situation", "start_sec", "end_sec", "duration_sec"]
    for c, header in enumerate(headers, 1):
        cell = chart_ws.cell(start_row, c, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")

    for i, (s0, s1, situation) in enumerate(spans, start_row + 1):
        chart_ws.cell(i, 1, situation)
        chart_ws.cell(i, 2, round(s0, 3))
        chart_ws.cell(i, 3, round(s1, 3))
        chart_ws.cell(i, 4, round(s1 - s0, 3))

    for col, width in {"A": 16, "B": 14, "C": 14, "D": 14}.items():
        chart_ws.column_dimensions[col].width = width


def parse_args():
    parser = argparse.ArgumentParser(description="集中度時系列グラフにシチュエーション背景を付けたシートを追加します。")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--png", type=Path, default=None, help="グラフ画像の保存先。省略時は output と同名の .png")
    parser.add_argument("--source-sheet", default="agg_timeseries_sec")
    parser.add_argument("--chart-sheet", default="状況背景グラフ")
    parser.add_argument("--situation-col", default="situation")
    parser.add_argument("--x-col", default="sec")
    parser.add_argument("--score-col", default="score_mean")
    parser.add_argument("--title", default="集中度の推移（背景：シチュエーション）")
    parser.add_argument("--width", type=int, default=1500)
    parser.add_argument("--height", type=int, default=820)
    return parser.parse_args()


def main():
    args = parse_args()
    png_path = args.png or args.output.with_suffix(".png")

    wb = openpyxl.load_workbook(args.input)
    if args.source_sheet not in wb.sheetnames:
        raise RuntimeError(f"シートが見つかりません: {args.source_sheet}")

    rows = collect_rows(
        wb[args.source_sheet],
        situation_col=args.situation_col,
        x_col=args.x_col,
        score_col=args.score_col,
    )
    spans, _dt = build_situation_spans(rows)

    make_chart_png(
        rows,
        spans,
        png_path,
        title=args.title,
        x_label=args.x_col,
        y_label=args.score_col,
        width=args.width,
        height=args.height,
    )
    add_chart_sheet(wb, png_path, spans, sheet_name=args.chart_sheet, title=args.title)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print("saved workbook:", args.output.resolve())
    print("saved chart image:", png_path.resolve())
    print(f"rows={len(rows)} spans={len(spans)}")


if __name__ == "__main__":
    main()
