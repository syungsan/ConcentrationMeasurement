from __future__ import annotations

import argparse
import json
import math
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.axis import ChartLines
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SITUATIONS = ("聞く", "書く", "話し合う")
PHASE_ORDER = ("pre", "AB", "cooldown", "post")
COLORS = {
    "navy": "1D3F66",
    "blue": "4472C4",
    "light_blue": "DCEBF7",
    "orange": "ED7D31",
    "green": "70AD47",
    "red": "C00000",
    "light_red": "FCE4D6",
    "gray": "E7E6E6",
    "white": "FFFFFF",
}


@dataclass(frozen=True)
class ABEvent:
    event_id: str
    video_path: str
    start_sec: float
    end_sec: float
    source: str


def parse_number_list(value: str) -> list[float]:
    values = sorted({float(item.strip()) for item in value.split(",") if item.strip()})
    if not values or any(item < 0 for item in values):
        raise ValueError("cooldown sensitivity values must be non-negative")
    return values


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create an AB intervention analysis workbook. AB is treated as an intervention "
            "state, while the original lesson situation is retained separately."
        )
    )
    parser.add_argument("--db", required=True, type=Path, help="pred_log.sqlite")
    parser.add_argument("--output", required=True, type=Path, help="output .xlsx")
    parser.add_argument("--video-filter", default="", help="Substring filter for prediction video_path")
    parser.add_argument("--ab-start", type=float, help="Explicit AB start in seconds")
    parser.add_argument("--ab-end", type=float, help="Explicit AB end in seconds")
    parser.add_argument("--pre-sec", type=float, default=300.0)
    parser.add_argument("--pre-gap-sec", type=float, default=0.0)
    parser.add_argument("--post-sec", type=float, default=300.0)
    parser.add_argument("--cooldown-sec", type=float, default=90.0)
    parser.add_argument("--cooldown-sensitivity", default="0,30,60,90,120,180")
    parser.add_argument("--bin-sec", type=float, default=5.0)
    parser.add_argument("--trim-proportion", type=float, default=0.10)
    parser.add_argument("--min-tracks", type=int, default=3)
    parser.add_argument("--bootstrap", type=int, default=10000)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    if (args.ab_start is None) != (args.ab_end is None):
        parser.error("--ab-start and --ab-end must be supplied together")
    if args.ab_start is not None and args.ab_end <= args.ab_start:
        parser.error("--ab-end must be greater than --ab-start")
    for name in ("pre_sec", "pre_gap_sec", "post_sec", "cooldown_sec"):
        if float(getattr(args, name)) < 0:
            parser.error(f"--{name.replace('_', '-')} must be non-negative")
    if args.bin_sec <= 0 or args.min_tracks < 1:
        parser.error("--bin-sec must be > 0 and --min-tracks must be >= 1")
    if not 0 <= args.trim_proportion < 0.5:
        parser.error("--trim-proportion must be in [0, 0.5)")
    return args


def table_names(conn: sqlite3.Connection) -> set[str]:
    return {str(row[0]) for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}


def table_columns(conn: sqlite3.Connection, table: str) -> set[str]:
    return {str(row[1]) for row in conn.execute(f"PRAGMA table_info({table})")}


def load_predictions(conn: sqlite3.Connection, video_filter: str) -> pd.DataFrame:
    columns = table_columns(conn, "predictions")
    score_col = "score" if "score" in columns else "yhat" if "yhat" in columns else None
    clamped_col = (
        "score_clamped" if "score_clamped" in columns
        else "score_int" if "score_int" in columns else None
    )
    video_col = "video_path" if "video_path" in columns else "video_id" if "video_id" in columns else None
    required = {"t", "track_id", "situation"}
    missing = sorted(required - columns)
    if missing or score_col is None or video_col is None:
        raise RuntimeError(f"predictions table lacks required columns: {missing}")
    override_sql = "situation_override" if "situation_override" in columns else "NULL"
    confidence_sql = "situation_confidence" if "situation_confidence" in columns else "NULL"
    clamped_sql = clamped_col if clamped_col else score_col
    query = f"""
        SELECT CAST({video_col} AS TEXT) AS video_path,
               CAST(t AS REAL) AS t,
               CAST(track_id AS INTEGER) AS track_id,
               CAST({score_col} AS REAL) AS score,
               CAST({clamped_sql} AS REAL) AS score_clamped,
               CAST(situation AS TEXT) AS base_situation,
               {override_sql} AS intervention_override,
               {confidence_sql} AS situation_confidence
        FROM predictions
    """
    frame = pd.read_sql_query(query, conn)
    if video_filter:
        frame = frame[frame["video_path"].str.contains(video_filter, na=False, regex=False)].copy()
    frame = frame.replace([np.inf, -np.inf], np.nan).dropna(subset=["t", "track_id", "score"])
    frame["base_situation"] = frame["base_situation"].astype(str)
    frame["is_ab"] = frame["intervention_override"].astype(str).eq("AB")
    return frame.sort_values(["video_path", "t", "track_id"]).reset_index(drop=True)


def infer_ab_runs(frame: pd.DataFrame) -> list[ABEvent]:
    events: list[ABEvent] = []
    for video_path, group in frame[frame["is_ab"]].groupby("video_path"):
        seconds = np.sort(group["t"].unique().astype(float))
        if len(seconds) == 0:
            continue
        gaps = np.diff(seconds)
        positive = gaps[gaps > 0]
        tolerance = max(2.0, float(np.median(positive) * 3.0) if len(positive) else 2.0)
        start = float(seconds[0])
        prior = float(seconds[0])
        run_number = 1
        for current in seconds[1:]:
            current = float(current)
            if current - prior > tolerance:
                events.append(ABEvent(f"AB{len(events)+1:03d}", str(video_path), start, prior + tolerance / 3.0, "prediction rows"))
                run_number += 1
                start = current
            prior = current
        events.append(ABEvent(f"AB{len(events)+1:03d}", str(video_path), start, prior + tolerance / 3.0, "prediction rows"))
    return events


def load_events(
    conn: sqlite3.Connection,
    frame: pd.DataFrame,
    *,
    video_filter: str,
    explicit_start: float | None,
    explicit_end: float | None,
) -> list[ABEvent]:
    videos = sorted(frame["video_path"].dropna().astype(str).unique())
    if explicit_start is not None and explicit_end is not None:
        return [
            ABEvent(f"AB{index:03d}", video, float(explicit_start), float(explicit_end), "CLI")
            for index, video in enumerate(videos, 1)
        ]
    events: list[ABEvent] = []
    if "situation_overrides" in table_names(conn):
        rows = conn.execute(
            """
            SELECT video_path, start_sec, end_sec
            FROM situation_overrides
            WHERE situation='AB'
            ORDER BY id
            """
        ).fetchall()
        keys: set[tuple[str, float, float]] = set()
        for raw_filter, start, end in rows:
            if start is None or end is None or float(end) <= float(start):
                continue
            applicable = videos
            if raw_filter:
                applicable = [video for video in videos if str(raw_filter) in video]
            for video in applicable:
                key = (video, float(start), float(end))
                if key not in keys:
                    keys.add(key)
                    events.append(ABEvent("", video, float(start), float(end), "situation_overrides"))
    if not events:
        events = infer_ab_runs(frame)
    events = sorted(events, key=lambda event: (event.video_path, event.start_sec, event.end_sec))
    return [
        ABEvent(f"AB{index:03d}", event.video_path, event.start_sec, event.end_sec, event.source)
        for index, event in enumerate(events, 1)
    ]


def trimmed_mean(values: Iterable[float], proportion: float) -> float:
    array = np.sort(np.asarray(list(values), dtype=float))
    array = array[np.isfinite(array)]
    if len(array) == 0:
        return float("nan")
    trim = int(math.floor(len(array) * proportion))
    if trim > 0 and trim * 2 < len(array):
        array = array[trim:-trim]
    return float(array.mean())


def assign_phase(
    time_value: pd.Series,
    event: ABEvent,
    *,
    pre_sec: float,
    pre_gap_sec: float,
    cooldown_sec: float,
    post_sec: float,
) -> pd.Series:
    pre_start = event.start_sec - pre_gap_sec - pre_sec
    pre_end = event.start_sec - pre_gap_sec
    conditions = [
        time_value.ge(pre_start) & time_value.lt(pre_end),
        time_value.ge(event.start_sec) & time_value.lt(event.end_sec),
        time_value.ge(event.end_sec) & time_value.lt(event.end_sec + cooldown_sec),
        time_value.ge(event.end_sec + cooldown_sec) & time_value.lt(event.end_sec + cooldown_sec + post_sec),
    ]
    return pd.Series(np.select(conditions, PHASE_ORDER, default="outside"), index=time_value.index)


def aggregate_event(
    frame: pd.DataFrame,
    event: ABEvent,
    *,
    pre_sec: float,
    pre_gap_sec: float,
    cooldown_sec: float,
    post_sec: float,
    bin_sec: float,
    trim_proportion: float,
    min_tracks: int,
) -> pd.DataFrame:
    data = frame[frame["video_path"] == event.video_path].copy()
    data["phase"] = assign_phase(
        data["t"], event,
        pre_sec=pre_sec, pre_gap_sec=pre_gap_sec,
        cooldown_sec=cooldown_sec, post_sec=post_sec,
    )
    data = data[data["phase"] != "outside"].copy()
    if data.empty:
        return pd.DataFrame()
    data["relative_to_ab_end_sec"] = data["t"] - event.end_sec
    data["relative_bin"] = np.floor(data["relative_to_ab_end_sec"] / bin_sec).astype(int)
    data["bin_start_relative_sec"] = data["relative_bin"] * bin_sec
    data["bin_center_relative_sec"] = data["bin_start_relative_sec"] + bin_sec / 2.0

    per_track = data.groupby(
        ["phase", "relative_bin", "bin_start_relative_sec", "bin_center_relative_sec", "base_situation", "track_id"],
        as_index=False,
    ).agg(
        score=("score", "mean"),
        score_clamped=("score_clamped", "mean"),
        situation_confidence=("situation_confidence", "mean"),
    )

    def summarize(group: pd.DataFrame, base_situation: str) -> dict[str, object]:
        scores = group["score"].to_numpy(dtype=float)
        return {
            "event_id": event.event_id,
            "video_path": event.video_path,
            "ab_start_sec": event.start_sec,
            "ab_end_sec": event.end_sec,
            "phase": str(group["phase"].iloc[0]),
            "relative_bin": int(group["relative_bin"].iloc[0]),
            "bin_start_relative_sec": float(group["bin_start_relative_sec"].iloc[0]),
            "bin_center_relative_sec": float(group["bin_center_relative_sec"].iloc[0]),
            "base_situation": base_situation,
            "n_tracks": int(group["track_id"].nunique()),
            "score_mean": float(np.mean(scores)),
            "score_trimmed_mean": trimmed_mean(scores, trim_proportion),
            "score_p50": float(np.median(scores)),
            "score_std": float(np.std(scores)),
            "low_score_ratio": float(np.mean(scores <= 3.0)),
            "situation_confidence": float(group["situation_confidence"].mean()),
            "clamped_mean": float(group["score_clamped"].mean()),
        }

    rows: list[dict[str, object]] = []
    bin_keys = ["phase", "relative_bin", "bin_start_relative_sec", "bin_center_relative_sec"]
    for _, group in per_track.groupby(bin_keys, sort=True):
        rows.append(summarize(group, "ALL"))
    for keys, group in per_track.groupby(bin_keys + ["base_situation"], sort=True):
        rows.append(summarize(group, str(keys[-1])))
    result = pd.DataFrame(rows)
    result["quality_ok"] = result["n_tracks"] >= int(min_tracks)
    return result.sort_values(["event_id", "bin_center_relative_sec", "base_situation"]).reset_index(drop=True)


def situation_summary(course: pd.DataFrame) -> pd.DataFrame:
    eligible = course[
        course["quality_ok"]
        & course["phase"].isin(["pre", "post"])
        & course["base_situation"].isin(SITUATIONS)
    ].copy()
    if eligible.empty:
        return pd.DataFrame(columns=[
            "event_id", "video_path", "phase", "base_situation", "n_bins",
            "mean_score", "median_score", "mean_tracks", "mean_situation_confidence",
        ])
    return eligible.groupby(
        ["event_id", "video_path", "phase", "base_situation"], as_index=False
    ).agg(
        n_bins=("score_trimmed_mean", "size"),
        mean_score=("score_trimmed_mean", "mean"),
        median_score=("score_trimmed_mean", "median"),
        mean_tracks=("n_tracks", "mean"),
        mean_situation_confidence=("situation_confidence", "mean"),
    )


def effect_table(course: pd.DataFrame, sit_summary: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for event_id, event_rows in course.groupby("event_id"):
        video_path = str(event_rows["video_path"].iloc[0])
        overall = event_rows[
            event_rows["quality_ok"]
            & event_rows["base_situation"].eq("ALL")
            & event_rows["phase"].isin(["pre", "post"])
        ]
        means = overall.groupby("phase")["score_trimmed_mean"].mean()
        pre_unadjusted = float(means.get("pre", np.nan))
        post_unadjusted = float(means.get("post", np.nan))

        event_sit = sit_summary[sit_summary["event_id"] == event_id]
        pre = event_sit[event_sit["phase"] == "pre"].set_index("base_situation")
        post = event_sit[event_sit["phase"] == "post"].set_index("base_situation")
        shared = [situation for situation in SITUATIONS if situation in pre.index and situation in post.index]
        total_pre_bins = float(pre["n_bins"].sum()) if not pre.empty else 0.0
        shared_pre_bins = float(pre.loc[shared, "n_bins"].sum()) if shared else 0.0
        coverage = shared_pre_bins / total_pre_bins if total_pre_bins > 0 else float("nan")
        if shared_pre_bins > 0:
            weights = pre.loc[shared, "n_bins"].astype(float) / shared_pre_bins
            standardized_pre = float((weights * pre.loc[shared, "mean_score"]).sum())
            standardized_post = float((weights * post.loc[shared, "mean_score"]).sum())
        else:
            standardized_pre = standardized_post = float("nan")
        rows.append({
            "event_id": event_id,
            "video_path": video_path,
            "n_pre_bins": int((overall["phase"] == "pre").sum()),
            "n_post_bins": int((overall["phase"] == "post").sum()),
            "pre_unadjusted": pre_unadjusted,
            "post_unadjusted": post_unadjusted,
            "effect_unadjusted": post_unadjusted - pre_unadjusted,
            "pre_standardized": standardized_pre,
            "post_standardized": standardized_post,
            "effect_standardized": standardized_post - standardized_pre,
            "shared_situations": ", ".join(shared),
            "pre_weight_coverage": coverage,
        })
    return pd.DataFrame(rows)


def bootstrap_mean_ci(values: Sequence[float], iterations: int, seed: int) -> tuple[float, float, float, int]:
    array = np.asarray(values, dtype=float)
    array = array[np.isfinite(array)]
    if len(array) == 0:
        return float("nan"), float("nan"), float("nan"), 0
    mean = float(array.mean())
    if len(array) < 2 or iterations <= 0:
        return mean, float("nan"), float("nan"), int(len(array))
    rng = np.random.default_rng(seed)
    samples = rng.choice(array, size=(iterations, len(array)), replace=True).mean(axis=1)
    return mean, float(np.quantile(samples, 0.025)), float(np.quantile(samples, 0.975)), int(len(array))


def sensitivity_analysis(
    frame: pd.DataFrame,
    events: list[ABEvent],
    cooldowns: list[float],
    args: argparse.Namespace,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    event_rows: list[pd.DataFrame] = []
    summary_rows: list[dict[str, object]] = []
    for cooldown in cooldowns:
        courses = [
            aggregate_event(
                frame, event,
                pre_sec=args.pre_sec, pre_gap_sec=args.pre_gap_sec,
                cooldown_sec=cooldown, post_sec=args.post_sec,
                bin_sec=args.bin_sec, trim_proportion=args.trim_proportion,
                min_tracks=args.min_tracks,
            )
            for event in events
        ]
        course = pd.concat([item for item in courses if not item.empty], ignore_index=True) if any(not item.empty for item in courses) else pd.DataFrame()
        effects = effect_table(course, situation_summary(course)) if not course.empty else pd.DataFrame()
        if not effects.empty:
            effects.insert(0, "cooldown_sec", cooldown)
            event_rows.append(effects)
            for metric in ("effect_unadjusted", "effect_standardized"):
                mean, low, high, n_events = bootstrap_mean_ci(
                    effects[metric].to_numpy(dtype=float), args.bootstrap, args.seed + int(cooldown * 10)
                )
                summary_rows.append({
                    "cooldown_sec": cooldown,
                    "effect_type": metric,
                    "n_events": n_events,
                    "mean_effect": mean,
                    "ci95_low": low,
                    "ci95_high": high,
                })
    return (
        pd.concat(event_rows, ignore_index=True) if event_rows else pd.DataFrame(),
        pd.DataFrame(summary_rows),
    )


def events_frame(events: list[ABEvent], args: argparse.Namespace) -> pd.DataFrame:
    return pd.DataFrame([{
        "event_id": event.event_id,
        "video_path": event.video_path,
        "ab_start_sec": event.start_sec,
        "ab_end_sec": event.end_sec,
        "ab_duration_sec": event.end_sec - event.start_sec,
        "pre_start_sec": event.start_sec - args.pre_gap_sec - args.pre_sec,
        "pre_end_sec": event.start_sec - args.pre_gap_sec,
        "cooldown_end_sec": event.end_sec + args.cooldown_sec,
        "post_end_sec": event.end_sec + args.cooldown_sec + args.post_sec,
        "event_source": event.source,
    } for event in events])


def aggregate_chart_data(course: pd.DataFrame, sit_summary: pd.DataFrame, effects: pd.DataFrame, sensitivity: pd.DataFrame) -> dict[str, pd.DataFrame]:
    overall_course_long = course[
        course["quality_ok"] & course["base_situation"].eq("ALL")
    ].groupby(["bin_center_relative_sec", "phase"], as_index=False).agg(
        score_trimmed_mean=("score_trimmed_mean", "mean"),
        n_tracks=("n_tracks", "mean"),
        n_events=("event_id", "nunique"),
    )
    overall_course = overall_course_long.groupby("bin_center_relative_sec", as_index=False).agg(
        score_trimmed_mean=("score_trimmed_mean", "mean")
    )
    situation_plot = sit_summary.groupby(["base_situation", "phase"], as_index=False).agg(
        mean_score=("mean_score", "mean")
    )
    situation_wide = situation_plot.pivot(index="base_situation", columns="phase", values="mean_score").reset_index()
    for phase in ("pre", "post"):
        if phase not in situation_wide:
            situation_wide[phase] = np.nan
    situation_wide = situation_wide[["base_situation", "pre", "post"]]
    effect_plot = pd.DataFrame({
        "effect_type": ["unadjusted", "situation-standardized"],
        "pre": [effects["pre_unadjusted"].mean(), effects["pre_standardized"].mean()],
        "post": [effects["post_unadjusted"].mean(), effects["post_standardized"].mean()],
        "effect": [effects["effect_unadjusted"].mean(), effects["effect_standardized"].mean()],
    }) if not effects.empty else pd.DataFrame(columns=["effect_type", "pre", "post", "effect"])
    cooldown_plot = sensitivity[sensitivity["effect_type"] == "effect_standardized"].copy()
    return {
        "time_course": overall_course,
        "situation": situation_wide,
        "effect": effect_plot,
        "cooldown": cooldown_plot,
    }


def safe_value(value: object) -> object:
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return None if not math.isfinite(float(value)) else float(value)
    if pd.isna(value):
        return None
    return value


def write_dataframe(ws, frame: pd.DataFrame, start_row: int = 1, start_col: int = 1) -> tuple[int, int]:
    for col_offset, column in enumerate(frame.columns):
        ws.cell(start_row, start_col + col_offset, str(column))
    for row_offset, row in enumerate(frame.itertuples(index=False, name=None), 1):
        for col_offset, value in enumerate(row):
            ws.cell(start_row + row_offset, start_col + col_offset, safe_value(value))
    return start_row + len(frame), start_col + len(frame.columns) - 1


def style_table(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[start_row]:
        if start_col <= cell.column <= end_col:
            cell.fill = PatternFill("solid", fgColor=COLORS["blue"])
            cell.font = Font(color=COLORS["white"], bold=True)
            cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=start_row + 1, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.border = Border(bottom=thin)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    ws.auto_filter.ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
    ws.freeze_panes = ws.cell(start_row + 1, start_col)


def fit_columns(ws, max_width: int = 42) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        width = min(max([len(value) for value in values], default=8) + 2, max_width)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(10, width)


def add_title(ws, title: str, subtitle: str = "") -> None:
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].fill = PatternFill("solid", fgColor=COLORS["navy"])
    ws["A1"].font = Font(color=COLORS["white"], bold=True, size=16)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    if subtitle:
        ws.merge_cells("A2:H2")
        ws["A2"] = subtitle
        ws["A2"].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[2].height = 34


def set_series_names(chart, names: list[str]) -> None:
    for series, name in zip(chart.series, names):
        series.tx = SeriesLabel(v=name)


def line_chart(
    ws,
    data_ws,
    category_col: int,
    data_min_col: int,
    data_max_col: int,
    min_row: int,
    max_row: int,
    title: str,
    anchor: str,
    series_names: list[str],
    series_colors: list[str],
    x_title: str,
    y_title: str,
    y_min: float | None = None,
    y_max: float | None = None,
    show_legend: bool = True,
    y_number_format: str | None = None,
) -> None:
    if max_row <= min_row:
        return
    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.height = 8.2
    chart.width = 15.5
    if show_legend:
        chart.legend.position = "b"
    else:
        chart.legend = None
    chart.y_axis.title = y_title
    chart.x_axis.title = x_title
    if y_min is not None:
        chart.y_axis.scaling.min = y_min
    if y_max is not None:
        chart.y_axis.scaling.max = y_max
    if y_number_format is not None:
        chart.y_axis.numFmt = y_number_format
    chart.y_axis.majorGridlines = ChartLines()
    data = Reference(data_ws, min_col=data_min_col, max_col=data_max_col, min_row=min_row, max_row=max_row)
    cats = Reference(data_ws, min_col=category_col, min_row=min_row + 1, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    set_series_names(chart, series_names)
    chart.display_blanks = "gap"
    for series, color in zip(chart.series, series_colors):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 22000
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def build_workbook(
    output: Path,
    args: argparse.Namespace,
    events: pd.DataFrame,
    course: pd.DataFrame,
    sit_summary: pd.DataFrame,
    effects: pd.DataFrame,
    sensitivity_events: pd.DataFrame,
    sensitivity_summary: pd.DataFrame,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    dashboard = wb.create_sheet("Dashboard")
    readme = wb.create_sheet("README")
    datasets = {
        "Events": events,
        "Time Course": course,
        "Situation Summary": sit_summary,
        "Effects": effects,
        "Cooldown Events": sensitivity_events,
        "Cooldown Summary": sensitivity_summary,
    }
    chart_data = aggregate_chart_data(course, sit_summary, effects, sensitivity_summary)

    add_title(dashboard, "AB介入効果ダッシュボード", "AB中は授業集中度から分離し、主要解析では除外しています。")
    dashboard.sheet_view.showGridLines = False
    settings = [
        ("主要cooldown", args.cooldown_sec), ("AB前区間（秒）", args.pre_sec),
        ("AB後区間（秒）", args.post_sec), ("集計幅（秒）", args.bin_sec),
        ("最低検出児童数", args.min_tracks), ("ABイベント数", len(events)),
    ]
    dashboard["A4"] = "設定"
    dashboard["A4"].fill = PatternFill("solid", fgColor=COLORS["blue"])
    dashboard["A4"].font = Font(color=COLORS["white"], bold=True)
    for index, (label, value) in enumerate(settings, 5):
        dashboard.cell(index, 1, label)
        dashboard.cell(index, 2, safe_value(value))

    dashboard["D4"] = "主要結果（イベント平均）"
    dashboard["D4"].fill = PatternFill("solid", fgColor=COLORS["blue"])
    dashboard["D4"].font = Font(color=COLORS["white"], bold=True)
    key_results = [
        ("未調整 前", effects["pre_unadjusted"].mean() if not effects.empty else np.nan),
        ("未調整 後", effects["post_unadjusted"].mean() if not effects.empty else np.nan),
        ("未調整 効果", effects["effect_unadjusted"].mean() if not effects.empty else np.nan),
        ("標準化 前", effects["pre_standardized"].mean() if not effects.empty else np.nan),
        ("標準化 後", effects["post_standardized"].mean() if not effects.empty else np.nan),
        ("標準化 効果", effects["effect_standardized"].mean() if not effects.empty else np.nan),
        ("共通状況カバー率", effects["pre_weight_coverage"].mean() if not effects.empty else np.nan),
    ]
    for index, (label, value) in enumerate(key_results, 5):
        dashboard.cell(index, 4, label)
        dashboard.cell(index, 5, safe_value(value))
        dashboard.cell(index, 5).number_format = "0.0000"
    dashboard["E11"].number_format = "0.0%"
    dashboard["A13"] = "解釈上の注意"
    dashboard["A13"].font = Font(bold=True, color=COLORS["red"])
    dashboard.merge_cells("A14:F17")
    dashboard["A14"] = (
        "AB中は立位・身体運動を含むため集中度効果の計算から除外しています。"
        "標準化効果はAB前のシチュエーション構成比をAB前後へ共通適用した差です。"
        "cooldownの長さで結論が変わらないか、Cooldown Summaryも必ず確認してください。"
        "1イベントだけの場合は記述的結果であり、因果効果の断定はできません。"
    )
    dashboard["A14"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard["A14"].fill = PatternFill("solid", fgColor=COLORS["light_red"])

    chart_ws = wb.create_sheet("Chart Data")
    chart_ws.sheet_state = "hidden"
    row_cursor = 1
    locations: dict[str, tuple[int, int, int, int]] = {}
    for name, frame in chart_data.items():
        chart_ws.cell(row_cursor, 1, name)
        start = row_cursor + 1
        end_row, end_col = write_dataframe(chart_ws, frame, start_row=start)
        locations[name] = (start, end_row, 1, end_col)
        row_cursor = end_row + 3

    start, end, _, _ = locations["time_course"]
    mean_ab_duration = events["ab_duration_sec"].mean() if not events.empty else 0.0
    line_chart(
        dashboard, chart_ws, 1, 2, 2, start, end,
        (f"集中度推移（AB終了=0秒）\n"
         f"AB前: ～-{mean_ab_duration:g}秒｜AB中: -{mean_ab_duration:g}～0秒｜"
         f"cooldown: 0～{args.cooldown_sec:g}秒｜AB後: {args.cooldown_sec:g}秒～"), "H3",
        ["5秒トリム平均"], ["4472C4"],
        "AB終了からの秒数", "集中度", 1, 7, False, "0.0",
    )

    start, end, _, _ = locations["effect"]
    if end > start:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "AB前後：未調整とシチュエーション標準化"
        chart.height = 8.2
        chart.width = 15.5
        chart.y_axis.title = "集中度"
        chart.y_axis.scaling.min = 1
        chart.y_axis.scaling.max = 7
        chart.add_data(Reference(chart_ws, min_col=2, max_col=3, min_row=start, max_row=end), titles_from_data=True)
        set_series_names(chart, ["AB前", "AB後"])
        chart.set_categories(Reference(chart_ws, min_col=1, min_row=start + 1, max_row=end))
        dashboard.add_chart(chart, "H20")

    start, end, _, _ = locations["situation"]
    if end > start:
        chart = BarChart()
        chart.type = "col"
        chart.style = 11
        chart.title = "シチュエーション別 AB前後"
        chart.height = 8.2
        chart.width = 15.5
        chart.y_axis.title = "集中度"
        chart.y_axis.scaling.min = 1
        chart.y_axis.scaling.max = 7
        chart.add_data(Reference(chart_ws, min_col=2, max_col=3, min_row=start, max_row=end), titles_from_data=True)
        set_series_names(chart, ["AB前", "AB後"])
        chart.set_categories(Reference(chart_ws, min_col=1, min_row=start + 1, max_row=end))
        dashboard.add_chart(chart, "H37")

    start, end, _, _ = locations["cooldown"]
    if end > start:
        line_chart(
            dashboard, chart_ws, 1, 4, 4, start, end,
            "クールダウン設定による推定AB効果の変化", "H54", ["標準化効果"], ["4472C4"],
            "AB終了後に除外する時間（秒）", "集中度差（AB後−AB前）",
            show_legend=False, y_number_format="0.00",
        )

    add_title(readme, "AB介入効果レポート：解析仕様")
    readme.sheet_view.showGridLines = False
    notes = [
        ("基本方針", "ABはシチュエーションではなく介入状態として扱います。元の聞く・書く・話し合うはbase_situationとして保持します。"),
        ("主要効果", "AB前のシチュエーション構成比を固定し、AB前後の各状況平均へ同じ重みを適用した標準化差です。"),
        ("AB中", "立位・運動のため、授業集中度の主要解析から除外します。Time Courseには参考値として残します。"),
        ("cooldown", f"主要設定はAB終了後{args.cooldown_sec:g}秒です。0～180秒の感度分析も出力します。"),
        ("時間集計", f"各児童内を{args.bin_sec:g}秒ごとに平均してから教室内を集約します。主要値は{args.trim_proportion:.0%}トリム平均です。"),
        ("品質条件", f"検出児童数が{args.min_tracks}人未満の時間ビンは主要集計から除外します。"),
        ("推奨確認", "Dashboard、Effects、Situation Summary、Cooldown Summary、Time Courseの順に確認してください。"),
        ("限界", "対照群のない単一授業の前後差は、時間経過・授業展開・測定誤差を完全には除けません。複数授業ではイベント単位CIを確認してください。"),
    ]
    for row, (label, text) in enumerate(notes, 4):
        readme.cell(row, 1, label).font = Font(bold=True, color=COLORS["navy"])
        readme.cell(row, 2, text).alignment = Alignment(wrap_text=True, vertical="top")
        readme.row_dimensions[row].height = 42
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 100

    for sheet_name, frame in datasets.items():
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.showGridLines = False
        add_title(ws, sheet_name)
        end_row, end_col = write_dataframe(ws, frame, start_row=3)
        if len(frame):
            style_table(ws, 3, end_row, 1, end_col)
        fit_columns(ws)
        if sheet_name in {"Effects", "Cooldown Summary"} and len(frame):
            numeric_start = 3
            numeric_end = end_row
            effect_columns = [
                index + 1 for index, column in enumerate(frame.columns)
                if "effect" in str(column) or str(column) == "mean_effect"
            ]
            for column in effect_columns:
                ws.conditional_formatting.add(
                    f"{get_column_letter(column)}{numeric_start+1}:{get_column_letter(column)}{numeric_end}",
                    ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
                )

    dashboard.column_dimensions["A"].width = 24
    dashboard.column_dimensions["B"].width = 14
    dashboard.column_dimensions["D"].width = 24
    dashboard.column_dimensions["E"].width = 14
    dashboard.freeze_panes = "A4"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> None:
    args = parse_args()
    db_path = args.db.resolve()
    if not db_path.exists():
        raise FileNotFoundError(db_path)
    cooldowns = parse_number_list(args.cooldown_sensitivity)
    if args.cooldown_sec not in cooldowns:
        cooldowns.append(float(args.cooldown_sec))
        cooldowns.sort()
    with sqlite3.connect(str(db_path)) as conn:
        frame = load_predictions(conn, args.video_filter)
        if frame.empty:
            raise RuntimeError("no prediction rows matched the requested video")
        events = load_events(
            conn, frame, video_filter=args.video_filter,
            explicit_start=args.ab_start, explicit_end=args.ab_end,
        )
    if not events:
        raise RuntimeError(
            "No AB interval found. Apply scripts/tools/apply_ab_interval.py first, "
            "or provide --ab-start and --ab-end."
        )
    primary_parts = [
        aggregate_event(
            frame, event,
            pre_sec=args.pre_sec, pre_gap_sec=args.pre_gap_sec,
            cooldown_sec=args.cooldown_sec, post_sec=args.post_sec,
            bin_sec=args.bin_sec, trim_proportion=args.trim_proportion,
            min_tracks=args.min_tracks,
        )
        for event in events
    ]
    primary_parts = [part for part in primary_parts if not part.empty]
    if not primary_parts:
        raise RuntimeError("AB intervals were found, but no rows fell inside the analysis windows")
    course = pd.concat(primary_parts, ignore_index=True)
    sit_summary = situation_summary(course)
    effects = effect_table(course, sit_summary)
    sensitivity_events, sensitivity_summary = sensitivity_analysis(frame, events, cooldowns, args)
    event_table = events_frame(events, args)
    output = args.output.resolve()
    build_workbook(
        output, args, event_table, course, sit_summary, effects,
        sensitivity_events, sensitivity_summary,
    )
    sidecar = output.with_suffix(".json")
    sidecar.write_text(json.dumps({
        "db": str(db_path),
        "output": str(output),
        "events": len(events),
        "settings": {
            "pre_sec": args.pre_sec,
            "pre_gap_sec": args.pre_gap_sec,
            "post_sec": args.post_sec,
            "cooldown_sec": args.cooldown_sec,
            "cooldown_sensitivity": cooldowns,
            "bin_sec": args.bin_sec,
            "trim_proportion": args.trim_proportion,
            "min_tracks": args.min_tracks,
        },
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Saved workbook: {output}")
    print(f"Saved run config: {sidecar}")
    print(effects.to_string(index=False))


if __name__ == "__main__":
    main()
